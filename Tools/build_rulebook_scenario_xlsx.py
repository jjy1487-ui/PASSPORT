# -*- coding: utf-8 -*-
"""규정집 시나리오 엑셀 생성.
- 일자별 규정 도입: 게임이 읽는 dayN.json의 rules(그 날 새로 추가되는 규정) 기준.
- 누적 규정집: 도입일차 순 전체 규정.
- 규정집 섹션(UI): 소스 엑셀 rule_section 시트(탭/항목/노출일자 등 구조화 규정집).
결과: data/규정집_시나리오.xlsx
"""
import json, io, os, glob, re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GAMEDATA = os.path.join(ROOT, "Assets", "Resources", "GameData")
SRC_XLSX = os.path.join(ROOT, "data", "여권_정리_updated.xlsx")
OUT = os.path.join(ROOT, "data", "규정집_시나리오.xlsx")

HDR_FILL = PatternFill("solid", fgColor="375623")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
NEW_FILL = PatternFill("solid", fgColor="E2EFDA")   # 신규 규정 도입일
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = BORDER
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

def setw(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

def border_wrap(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER; c.alignment = WRAP

# ── load dayN rules ────────────────────────────────────────────────
days = []
for p in sorted(glob.glob(os.path.join(GAMEDATA, "day*.json")),
                key=lambda p: int(re.search(r"day(\d+)", p).group(1))):
    dn = int(re.search(r"day(\d+)", p).group(1))
    d = json.load(io.open(p, encoding="utf-8"))
    days.append((dn, d.get("rules", []) or []))

wb = openpyxl.Workbook()

# ── Sheet: 일자별_규정도입 ──────────────────────────────────────────
ws = wb.active; ws.title = "일자별_규정도입"
cols = ["일차", "신규규정수", "규정ID", "규정 제목", "규정 내용", "연계필드"]
ws.append(cols)
for dn, rules in days:
    if not rules:
        ws.append([dn, 0, "", "(신규 규정 없음 — 규정집 변동 없음)", "", ""])
        continue
    for j, r in enumerate(rules):
        ws.append([dn, len(rules) if j == 0 else "", r.get("ruleId"),
                   r.get("title"), r.get("content"), r.get("relatedField", "") or ""])
        for c in range(1, len(cols)+1):
            ws.cell(row=ws.max_row, column=c).fill = NEW_FILL
style_header(ws, len(cols)); border_wrap(ws); setw(ws, [5, 9, 7, 18, 70, 12])

# ── Sheet: 누적_규정집 (도입일차 순) ────────────────────────────────
ws = wb.create_sheet("누적_규정집")
cols = ["도입일차", "규정ID", "규정 제목", "규정 내용"]
ws.append(cols)
cum = []
for dn, rules in days:
    for r in rules:
        cum.append((dn, r))
for dn, r in cum:
    ws.append([dn, r.get("ruleId"), r.get("title"), r.get("content")])
style_header(ws, len(cols)); border_wrap(ws); setw(ws, [9, 7, 20, 80])

# ── Sheet: 규정집_섹션(UI) from source xlsx rule_section ────────────
ws = wb.create_sheet("규정집_섹션(UI)")
secwb = openpyxl.load_workbook(SRC_XLSX, data_only=True)
sec = secwb["rule_section"]
secrows = list(sec.iter_rows(values_only=True))
# header is row index 2 (0-based) per schema dump: section_id,tab,tab_order,category,...
hdr = ["섹션ID","탭","탭순서","항목","항목순서","제목","본문","이미지키","내용유형","노출일자","우선순위","연계필드","비고"]
ws.append(hdr)
for r in secrows[4:]:  # data starts after PK/type/eng/kor header rows
    if r and r[0] not in (None, "", "section_id") and isinstance(r[0], (int, float)):
        ws.append([("" if v is None else v) for v in r[:13]])
style_header(ws, len(hdr)); border_wrap(ws)
setw(ws, [7, 8, 7, 12, 8, 18, 60, 12, 9, 8, 8, 10, 24])

# ── Sheet 0: 안내 ───────────────────────────────────────────────────
info = wb.create_sheet("0_안내", 0)
info["A1"] = "여권 주세요 — 규정집 시나리오"; info["A1"].font = Font(bold=True, size=14)
intro_days = [f"day{dn}: " + ", ".join(r.get("title") for r in rules) for dn, rules in days if rules]
lines = [
    "",
    "소스: 게임이 읽는 Assets/Resources/GameData/dayN.json 의 rules(그 날 새로 도입되는 규정)",
    "      + 소스 엑셀 data/여권_정리_updated.xlsx 의 rule_section(구조화 규정집 UI).",
    "",
    "■ 시트 구성",
    "  1) 일자별_규정도입 : 날짜별로 새로 추가되는 규정(초록색). 변동 없는 날도 표시.",
    "  2) 누적_규정집     : 도입일차 순 전체 규정(=14일차 시점 규정집 전체).",
    "  3) 규정집_섹션(UI) : 탭/항목/노출일자로 구조화된 규정집 UI 섹션.",
    "",
    "■ 규정 도입 흐름(시나리오)",
] + ["    " + s for s in intro_days] + [
    "",
    "※ 참고: dayN.json의 rules는 '그날 신규' 규정만 담겨 있습니다(누적 아님).",
    "   게임 규정집 팝업이 '누적'으로 보여야 한다면 별도 처리가 필요합니다(현재는 그날 rules만 Open).",
]
for i, t in enumerate(lines, start=2):
    info[f"A{i}"] = t
info.column_dimensions["A"].width = 105

wb.save(OUT)
print("saved:", OUT)
print("total rules:", len(cum), "| days with new rules:", sum(1 for _, r in days if r))
