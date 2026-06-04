# -*- coding: utf-8 -*-
"""소스 엑셀 3종의 시트/헤더/샘플행을 덤프한다. PYTHONUTF8=1로 실행."""
import sys, os
import openpyxl

FILES = {
    "schedule": r"C:\Users\chris\Downloads\dayeon_data\여권주세요_날짜별_방문고객_랜덤정리_3.xlsx",
    "script":   r"C:\Users\chris\Downloads\dayeon_data\방문객_스크립트_전체_260604.xlsx",
    "master":   r"C:\Users\chris\Downloads\여권_정리_updated.xlsx",
}

def dump(path, maxrows=8, maxcols=20):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        print("=" * 80)
        print(f"SHEET: {ws.title}  (dims={ws.max_row}x{ws.max_column})")
        print("-" * 80)
        for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if r > maxrows:
                print(f"... ({ws.max_row} rows total)")
                break
            cells = ["" if c is None else str(c) for c in row[:maxcols]]
            print(f"r{r}: " + " | ".join(cells))
    wb.close()

if __name__ == "__main__":
    which = sys.argv[1] if len(sys.argv) > 1 else "all"
    for key, path in FILES.items():
        if which != "all" and which != key:
            continue
        print("#" * 90)
        print(f"# FILE [{key}]: {path}")
        print("#" * 90)
        dump(path)
