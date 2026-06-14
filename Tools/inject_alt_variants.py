# -*- coding: utf-8 -*-
"""확률 캐릭터(박철수 등)에 '반대 변형'(altVariant) + validChance 를 주입한다.
- 신원(이름·얼굴·생년월일)은 baked 그대로, 서류 정상/불량(성형은 지문 본인/도용)만 반대로 만든 변형.
- 대사는 손글 원본의 반대 변형 블록에서 가져온다.
- 런타임 CustomerRoster.RollVariants 가 매 플레이 validChance 로 굴려 오버레이.
패턴: 진상→여권 만료일 / 일반→생년월일(정보불일치) / 외국인→사진 / 성형→지문(본인↔도용).
"""
import openpyxl, json, io, glob, re, copy

ROOT = r"C:\Users\chris\Documents\produc_build_reecture"
HAND = r"C:\Users\chris\Downloads\방문객_스크립트_일자별_생성_손글기반_260611.xlsx"
SRC_XLSX = ROOT + r"\data\여권_정리_updated.xlsx"
GAMEDATA = ROOT + r"\Assets\Resources\GameData"

APPROVE = "정상 승인"; REJECT = "정상 거절"
STD_REJECT_INSP = "확인되지 않은 사유로 입국이 어렵습니다. 죄송합니다."
DOC_FAIL_INSP = "서류 오류가 확인되어 입국하실 수 없습니다."
APPROVE_INSP = "즐거운 여행 되십시오."
STOLEN = [("박도윤","1995-02-18"),("강민호","1994-03-09"),("윤하늘","1996-12-01")]
# 방역 구간(5~7일) 진상 중 '불량 = PCR 미지참'으로 처리할 손님 customer_id.
# 조지호(47)만 지정. 홍성민(day9 진상)·오은우(day6)·신유준(day7)은 기존 경로(여권 만료) 유지 — 여기 넣지 않는다.
PCR_MISSING_CIDS = {47}
# 타입 기본 결함(진상→만료일/일반→생년월일)이 설계 대사와 어긋나는 손님: 결함 필드를 강제 지정.
#   박철수(5)=성별, 최서연(6)=만료일. make_defect 가 이 값을 보고 해당 필드 결함을 만든다.
DEFECT_FIELD_OVERRIDE = {5: "gender", 6: "expiry_date"}

# ── valid_chance per (day,slot) ──
wb = openpyxl.load_workbook(SRC_XLSX, data_only=True)
vc_map = {}
for r in list(wb["day_schedule"].iter_rows(values_only=True))[4:]:
    if r[1] is None or r[2] is None: continue
    try: vc_map[(int(float(r[1])), int(float(r[2])))] = float(r[4])
    except: pass

# ── hand-script blocks: (day,name)->[{docstate,intro,approve,reject}] ──
def clean(h):
    if h is None: return None
    m = re.search(r"「(.*?)」", str(h), re.S)
    return (m.group(1).strip() if m else str(h).split("※")[0].strip())
hw = openpyxl.load_workbook(HAND, data_only=True)
hand = {}
DSTEP = {"인삿말","검문관","방문객반응"}
for s in [x for x in hw.sheetnames if x.startswith("Day")]:
    day = int(re.search(r"(\d+)", s).group(1)); cur = None
    for r in hw[s].iter_rows(values_only=True):
        cells = ["" if c is None else str(c) for c in r]
        joined = " ".join(cells)
        if "번째 방문객]" in joined:
            hdr = next((c for c in cells if "번째 방문객]" in c), joined)
            m = re.search(r"\]\s*([^|]+?)\s*\|\s*([^|]+?)\s*\|\s*서류:\s*([^|]+)", hdr)
            if m:
                ds = re.sub(r"\(더미.*?\)", "", m.group(3)).strip()
                cur = {"docstate": ds, "name": m.group(1).strip(), "intro": [], "approve": [], "reject": []}
                hand.setdefault((day, m.group(1).strip()), []).append(cur)
            continue
        if cur is None: continue
        step = cells[1].strip() if len(cells) > 1 else ""
        if step not in DSTEP: continue
        result = cells[6].strip() if len(cells) > 6 else ""
        line = clean(cells[7] if len(cells) > 7 else "")
        if not line: continue
        spk = "심사관" if step == "검문관" else (cells[3].strip() or cur["name"])
        if step == "인삿말" or result == "공통": cur["intro"].append((spk, line))
        elif result == "허가": cur["approve"].append((spk, line))
        elif result == "거부": cur["reject"].append((spk, line))

def pick_block(day, name, want_normal):
    for b in hand.get((day, name), []):
        isn = ("정상" in b["docstate"]) and ("불량" not in b["docstate"])
        if isn == want_normal: return b
    return None

def insp_of(pairs):  return next((t for s,t in pairs if s == "심사관"), None)
def vis_of(pairs):   return next((t for s,t in pairs if s != "심사관"), None)
def lines(*pairs):   return [{"order": i+1, "speaker": s, "text": t} for i,(s,t) in enumerate(pairs)]

def build_cases(block, name, correct):
    intro = block["intro"] if block else []
    a_insp = (insp_of(block["approve"]) if block else None) or APPROVE_INSP
    a_vis  = (vis_of(block["approve"]) if block else None) or "감사합니다."
    r_insp = (insp_of(block["reject"]) if block else None) or (STD_REJECT_INSP if correct==APPROVE else DOC_FAIL_INSP)
    r_vis  = (vis_of(block["reject"]) if block else None) or "알겠습니다."
    intro_lines = [{"order": i+1, "speaker": s, "text": t} for i,(s,t) in enumerate(intro)] \
                  or [{"order":1,"speaker":name,"text":"안녕하세요."}]
    cases = [{"caseType":"입장","gameResult":"-","rejectCount":0,"lines":intro_lines}]
    if correct == APPROVE:  # alt 정상: 허가=정답, 거부=오판
        cases += [
            {"caseType":"일반 심사","gameResult":APPROVE,"rejectCount":0,"lines":lines(("심사관",a_insp),(name,a_vis))},
            {"caseType":"일반 심사","gameResult":REJECT,"rejectCount":0,"lines":lines(("심사관",DOC_FAIL_INSP),(name,r_vis))},
            {"caseType":"일반 심사","gameResult":"잘못 허가","rejectCount":0,"lines":lines(("심사관",a_insp),(name,a_vis))},
            {"caseType":"일반 심사","gameResult":"잘못 거절","rejectCount":0,"lines":lines(("심사관",STD_REJECT_INSP),(name,r_vis))},
        ]
    else:  # alt 불량: 거부=정답(구체 사유), 허가=오판
        cases += [
            {"caseType":"일반 심사","gameResult":APPROVE,"rejectCount":0,"lines":lines(("심사관",a_insp),(name,a_vis))},
            {"caseType":"일반 심사","gameResult":REJECT,"rejectCount":0,"lines":lines(("심사관",r_insp),(name,r_vis))},
            {"caseType":"일반 심사","gameResult":"잘못 허가","rejectCount":0,"lines":lines(("심사관",a_insp),(name,a_vis))},
            {"caseType":"일반 심사","gameResult":"잘못 거절","rejectCount":0,"lines":lines(("심사관",STD_REJECT_INSP),(name,r_vis))},
        ]
    return cases

def find_passport(docs):
    for d in docs:
        if d.get("documentType") == "여권": return d
    return docs[0] if docs else None
def set_field(doc, key, val):
    for f in doc.get("fields", []):
        if f.get("key") == key: f["value"] = val; return
def get_field(doc, key):
    for f in doc.get("fields", []):
        if f.get("key") == key: return f.get("value")
    return None

# sprite pool for 사진 불일치
sprite_pool = set()
for p in glob.glob(GAMEDATA + r"\day*.json"):
    if "backup" in p: continue
    for c in json.load(io.open(p, encoding="utf-8")).get("customers", []):
        if c.get("spriteRef"): sprite_pool.add(c["spriteRef"])
sprite_pool = sorted(sprite_pool)

def make_defect(docs, ctype, customer):
    docs = copy.deepcopy(docs)
    pp = find_passport(docs)
    if pp is None: return docs
    # 설계 대사와 맞추기 위한 결함 필드 강제 지정(타입 기본보다 우선).
    forced = DEFECT_FIELD_OVERRIDE.get(customer.get("customerId"))
    if forced == "gender":                    # 성별 불일치
        pp["variant"]="비정상"; pp["violationField"]="성별"
        g = get_field(pp, "gender")
        set_field(pp, "gender", "여성" if g == "남성" else "남성")
        return docs
    if forced == "expiry_date":               # 만료일 경과
        pp["variant"]="비정상"; pp["violationField"]="만료일"
        ex = get_field(pp,"expiry_date") or "2029-01-01"
        m = re.match(r"(\d{4})(-\d{2}-\d{2})", ex)
        if m: set_field(pp,"expiry_date", str(int(m.group(1))-5)+m.group(2))
        return docs
    if "진상" in ctype:                       # 여권 기간 오류
        pp["variant"]="비정상"; pp["violationField"]="만료일"
        ex = get_field(pp,"expiry_date") or "2029-01-01"
        m = re.match(r"(\d{4})(-\d{2}-\d{2})", ex)
        if m: set_field(pp,"expiry_date", str(int(m.group(1))-5)+m.group(2))
    elif "일반" in ctype:                      # 생년월일(정보 불일치)
        pp["variant"]="비정상"; pp["violationField"]="생년월일"; set_field(pp,"birth_date","1900-01-01")
    else:                                       # 외국인 → 사진 불일치
        pp["variant"]="비정상"; pp["violationField"]="사진"
        own = customer.get("spriteRef")
        alt_sprite = next((s for s in sprite_pool if s != own), own)
        pp["spriteRef"] = alt_sprite
    return docs

def make_valid(docs, customer):
    docs = copy.deepcopy(docs)
    pp = find_passport(docs)
    if pp is None: return docs
    pp["variant"]="정상"; pp["violationField"]="없음"
    pp["spriteRef"]=customer.get("spriteRef")
    if customer.get("nameEn"): set_field(pp,"name",customer["nameEn"])
    if customer.get("birthDate"): set_field(pp,"birth_date",customer["birthDate"])
    iss = get_field(pp,"issue_date") or "2019-01-01"
    m = re.match(r"\d{4}(-\d{2}-\d{2})", iss)
    if m: set_field(pp,"issue_date","2019"+m.group(1)); set_field(pp,"expiry_date","2029"+m.group(1))
    return docs

def make_pcr_missing(docs, customer):
    """PCR 미지참 변형: 여권은 정상으로 두고, PCR검사서 문서를 아예 제거(안 가져옴)."""
    docs = copy.deepcopy(docs)
    pp = find_passport(docs)
    if pp is not None:
        pp["variant"] = "정상"; pp["violationField"] = "없음"
        pp["spriteRef"] = customer.get("spriteRef")
    return [d for d in docs if d.get("documentType") != "PCR검사서"]

def apply_pcr_missing_dialogue(cases, name):
    """build_cases 결과의 입장/거절 라인을 PCR 미지참용으로 교체(day5.json 조지호와 동일)."""
    cases[0]["lines"] = [
        {"order": 1, "speaker": name, "text": "건강 이상자 때문에 심사 강화됐다는 거 알겠는데, 저는 멀쩡하다고요.", "claim": None},
        {"order": 2, "speaker": "심사관", "text": "오늘부터 방역 절차가 시행됩니다. 여권과 PCR 검사서를 함께 보여주세요.", "claim": None},
        {"order": 3, "speaker": name, "text": "어? PCR요? 그건 안 가져왔는데요.", "claim": None},
    ]
    for case in cases:
        if case["gameResult"] == REJECT:
            case["lines"] = [
                {"order": 1, "speaker": "심사관", "text": "방역 대상 입국자가 PCR 검사서를 제출하지 않아 입국하실 수 없습니다.", "claim": None},
                {"order": 2, "speaker": name, "text": "아니 그걸 꼭 가져와야 돼요? 난 멀쩡하다니까!", "claim": None},
            ]
    return cases

def fp_self(customer):
    nm = customer.get("nameEn") or customer.get("nameKr")
    return {"type":"fingerprint","result":"일치","detail":"본인 일치","extra":nm,
            "claim":{"attr":"name","value":nm,"label":"지문 대조 신원","unlocksScan":""},
            "record":{"mode":"성형","dbName":nm,"dbBirth":customer.get("birthDate",""),
                      "dbNationality":"KOR","criminalRecord":"없음","wantedNo":""}}
def fp_stolen(customer):
    nm, bd = STOLEN[customer.get("customerId",0) % len(STOLEN)]
    return {"type":"fingerprint","result":"불일치","detail":"신원 불일치","extra":nm,
            "claim":{"attr":"name","value":nm,"label":"지문 대조 신원","unlocksScan":""},
            "record":{"mode":"성형","dbName":nm,"dbBirth":bd,
                      "dbNationality":"KOR","criminalRecord":"없음","wantedNo":""}}

# ── inject ──
injected = 0; skipped = []
for p in sorted(glob.glob(GAMEDATA + r"\day*.json"), key=lambda p:int(re.search(r"day(\d+)",p).group(1))):
    if "backup" in p: continue
    day = int(re.search(r"day(\d+)", p).group(1))
    d = json.load(io.open(p, encoding="utf-8"))
    changed = False
    for c in d.get("customers", []):
        vc = vc_map.get((day, c.get("slot")))
        if vc is None or vc <= 0 or vc >= 1: continue
        name = c.get("nameKr"); ctype = c.get("characterType") or ""
        if name == "윤서린" or "수배" in ctype:   # 항상 수배자 → 변형 없음
            skipped.append((day,name)); continue
        base_normal = (c.get("correctResult") == APPROVE)
        alt_correct = REJECT if base_normal else APPROVE
        block = pick_block(day, name, want_normal=(alt_correct==APPROVE))
        # documents / fingerprint per pattern
        pcr_missing = (c.get("customerId") in PCR_MISSING_CIDS) and (alt_correct == REJECT)
        if "성형" in ctype:
            alt_docs = copy.deepcopy(c.get("documents", []))
            alt_fp = fp_self(c) if alt_correct==APPROVE else fp_stolen(c)
        else:
            alt_fp = c.get("fingerprint")
            if alt_correct == APPROVE:
                alt_docs = make_valid(c.get("documents",[]), c)
            elif pcr_missing:
                alt_docs = make_pcr_missing(c.get("documents",[]), c)
            else:
                alt_docs = make_defect(c.get("documents",[]), ctype, c)
        cases = build_cases(block, name, alt_correct)
        if pcr_missing:
            cases = apply_pcr_missing_dialogue(cases, name)
        c["validChance"] = round(vc, 3)
        c["altVariant"] = {
            "correctResult": alt_correct, "defectVariant": ("PCR 미제출" if pcr_missing else ""),
            "rejectAdvancedBranchKey": "", "rejectGuidedCaseType": "",
            "documents": alt_docs,
            "dialogueCases": cases,
            "xray": c.get("xray"), "fingerprint": alt_fp,
        }
        injected += 1; changed = True
    if changed:
        json.dump(d, io.open(p,"w",encoding="utf-8"), ensure_ascii=False, indent=2)
print(f"injected altVariant into {injected} customers")
print(f"skipped (always-수배자): {skipped}")
