# -*- coding: utf-8 -*-
"""dayN.json(런타임 시나리오 소스)에서 일자별 손님 큐 / 뉴스·경보 / 대사 스크립트를
한 권의 엑셀로 정리한다.  결과: data/시나리오_스크립트.xlsx

대사_스크립트: 손글 스크립트 형식 — 방문객별로 인삿말 + 의미 있는 2분기(입국 허가/입국 거부)만,
대사만(모션·UI 없음), 정답/오판 라벨. 의미 없는 중복 케이스(정상서류의 '잘못 허가' 등)와
오거부횟수 열은 넣지 않는다.
"""
import json, io, os, glob, re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GAMEDATA = os.path.join(ROOT, "Assets", "Resources", "GameData")
OUT = os.path.join(ROOT, "data", "시나리오_스크립트.xlsx")

# ── styles ──────────────────────────────────────────────────────────
HDR_FILL = PatternFill("solid", fgColor="2F5597")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
REJECT_FILL = PatternFill("solid", fgColor="FFF2CC")
ALERT_FILL = PatternFill("solid", fgColor="FCE4D6")
VISITOR_FILL = PatternFill("solid", fgColor="DDEBF7")   # 방문객 헤더
GROUP_FILL = PatternFill("solid", fgColor="F2F7FC")     # 짝수 방문객 음영
APPROVE_FONT = Font(color="1F7A1F", bold=True)          # 입국 허가 라벨
REJECT_FONT = Font(color="B02418", bold=True)           # 입국 거부 라벨
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

def style_header(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HDR_FILL; cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = BORDER
    ws.freeze_panes = "A2"

def setw(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

# ── load ────────────────────────────────────────────────────────────
days = []
for p in sorted(glob.glob(os.path.join(GAMEDATA, "day*.json")),
                key=lambda p: int(re.search(r"day(\d+)", p).group(1))):
    dn = int(re.search(r"day(\d+)", p).group(1))
    days.append((dn, json.load(io.open(p, encoding="utf-8"))))

def case_by_result(c, gr):
    for dc in c.get("dialogueCases", []):
        if dc.get("gameResult") == gr:
            return dc
    return None

def case_by_type(c, ct):
    if not ct:
        return None
    for dc in c.get("dialogueCases", []):
        if dc.get("caseType") == ct:
            return dc
    return None

def lines_sorted(case):
    if not case:
        return []
    return sorted(case.get("lines", []), key=lambda l: l.get("order", 0))

def defect_label(c):
    parts = []
    for doc in c.get("documents", []) or []:
        vf = doc.get("violationField", "")
        if (doc.get("variant") == "비정상") or (vf and vf != "없음"):
            parts.append(f"{doc.get('documentType','')} {vf}".strip())
    if c.get("xray") and c["xray"].get("detail"):
        parts.append(f"X-ray:{c['xray']['detail']}")
    fp = c.get("fingerprint")
    if fp and fp.get("record", {}).get("mode") == "수배자":
        parts.append("지문:수배자")
    return ", ".join(parts) if parts else "결함"

wb = openpyxl.Workbook()

# ── Sheet 1: 일자별_손님큐 ───────────────────────────────────────────
ws = wb.active; ws.title = "일자별_손님큐"
cols = ["일차","슬롯","고객ID","이름(한글)","이름(영문)","국적","성별","생년월일",
        "캐릭터유형","정답판정","결함변형","X-ray","지문","특수분기키","비고"]
ws.append(cols)
for dn, d in days:
    for c in d.get("customers", []):
        note = []
        if c.get("xray"): note.append("X-ray검사 보유")
        if c.get("fingerprint"): note.append("지문검사 보유")
        ws.append([dn, c.get("slot"), c.get("customerId"), c.get("nameKr"), c.get("nameEn"),
                   c.get("nationality"), c.get("gender"), c.get("birthDate"),
                   c.get("characterType"), c.get("correctResult"), c.get("defectVariant") or "",
                   "Y" if c.get("xray") else "", "Y" if c.get("fingerprint") else "",
                   c.get("rejectAdvancedBranchKey") or "", " / ".join(note)])
        r = ws.max_row
        if "거절" in (c.get("correctResult") or ""):
            for cc in range(1, len(cols)+1): ws.cell(row=r, column=cc).fill = REJECT_FILL
        for cc in range(1, len(cols)+1):
            ws.cell(row=r, column=cc).border = BORDER; ws.cell(row=r, column=cc).alignment = TOP
style_header(ws, len(cols)); ws.auto_filter.ref = ws.dimensions
setw(ws, [5,5,7,12,16,14,6,12,16,10,12,7,5,24,18])

# ── Sheet 2: 뉴스·경보 ──────────────────────────────────────────────
ws = wb.create_sheet("뉴스·경보")
cols = ["일차","뉴스ID","제목","내용","단서(attr)","단서(value)","잠금해제스캔","유형"]
ws.append(cols)
for dn, d in days:
    for n in d.get("news", []):
        claims = n.get("claims") or []
        c0 = claims[0] if claims else None
        kind = "경보(대조단서)" if c0 and c0.get("unlocksScan") else ("일반뉴스" if not c0 else "단서")
        ws.append([dn, n.get("newsId"), n.get("title"), n.get("content"),
                   c0.get("attr") if c0 else "", c0.get("value") if c0 else "",
                   c0.get("unlocksScan") if c0 else "", kind])
        r = ws.max_row
        if c0 and c0.get("unlocksScan"):
            for cc in range(1, len(cols)+1): ws.cell(row=r, column=cc).fill = ALERT_FILL
        for cc in range(1, len(cols)+1):
            ws.cell(row=r, column=cc).border = BORDER; ws.cell(row=r, column=cc).alignment = WRAP
style_header(ws, len(cols)); ws.auto_filter.ref = ws.dimensions
setw(ws, [5,9,30,70,12,16,14,14])

# ── Sheet 3: 대사_스크립트 (손글 형식: 인삿말 + 입국 허가/거부 2분기) ──
ws = wb.create_sheet("대사_스크립트")
cols = ["일차","순서","방문객","유형","서류상태","분기","화자","대사"]
ws.append(cols)
visitor_idx = 0
for dn, d in days:
    for c in d.get("customers", []):
        visitor_idx += 1
        name = c.get("nameKr"); slot = c.get("slot"); ctype = c.get("characterType")
        should_approve = (c.get("correctResult") == "정상 승인")
        docstate = "정상" if should_approve else f"불량 ({defect_label(c)})"
        shade = GROUP_FILL if (visitor_idx % 2 == 0) else None
        first_row_of_visitor = ws.max_row + 1

        blocks = []  # (분기라벨, 분기폰트, [lines])
        # 인삿말(입장)
        blocks.append(("인삿말", None, lines_sorted(case_by_type(c, "입장"))))
        # 입국 허가
        if should_approve:
            blocks.append(("입국 허가 (정답)", APPROVE_FONT, lines_sorted(case_by_result(c, "정상 승인"))))
        else:
            blocks.append(("입국 허가 (오판)", REJECT_FONT, lines_sorted(case_by_result(c, "잘못 허가"))))
        # 입국 거부
        if should_approve:
            blocks.append(("입국 거부 (오판)", REJECT_FONT, lines_sorted(case_by_result(c, "잘못 거절"))))
        else:
            guided = case_by_type(c, c.get("rejectGuidedCaseType"))
            if guided:
                blocks.append(("입국 거부 (정답·적발)", APPROVE_FONT, lines_sorted(guided)))
            else:
                blocks.append(("입국 거부 (정답)", APPROVE_FONT, lines_sorted(case_by_result(c, "정상 거절"))))

        for label, lfont, lines in blocks:
            for li, ln in enumerate(lines):
                ws.append([dn if li == 0 and label == "인삿말" else "",
                           slot if li == 0 and label == "인삿말" else "",
                           name if li == 0 and label == "인삿말" else "",
                           ctype if li == 0 and label == "인삿말" else "",
                           docstate if li == 0 and label == "인삿말" else "",
                           label if li == 0 else "",
                           ln.get("speaker"), ln.get("text")])
                r = ws.max_row
                if shade:
                    for cc in range(1, len(cols)+1): ws.cell(row=r, column=cc).fill = shade
                if li == 0 and lfont is not None:
                    ws.cell(row=r, column=6).font = lfont
                for cc in range(1, len(cols)+1):
                    ws.cell(row=r, column=cc).border = BORDER; ws.cell(row=r, column=cc).alignment = WRAP
        # 방문객 첫 줄 강조(이름 볼드)
        for cc in (3,):
            ws.cell(row=first_row_of_visitor, column=cc).font = Font(bold=True)
style_header(ws, len(cols))
setw(ws, [5,5,13,16,26,16,10,60])

# ── Sheet 0(맨앞): 안내 ─────────────────────────────────────────────
info = wb.create_sheet("0_안내", 0)
info["A1"] = "여권 주세요 — 현재 시나리오 스크립트"; info["A1"].font = Font(bold=True, size=14)
lines = [
    "",
    "소스: Assets/Resources/GameData/dayN.json (게임이 실제로 읽는 런타임 시나리오)",
    f"포함 일차: day1 ~ day{days[-1][0]}",
    "",
    "■ 시트 구성",
    "  1) 일자별_손님큐 : 날짜별 등장 순서/유형/정답판정/결함/특수검사 (노란색=정답이 '거절')",
    "  2) 뉴스·경보      : 날짜별 뉴스(주황색=대조 단서 있는 경보)",
    "  3) 대사_스크립트  : 방문객별 [인삿말 + 입국 허가 + 입국 거부] 대사만. (손글 형식)",
    "",
    "■ 대사_스크립트 읽는 법",
    "  - 손님마다 의미 있는 2분기만 표시: 입국 허가 / 입국 거부.",
    "  - (정답)=옳은 판정, (오판)=틀린 판정. 정상서류는 허가가 정답, 불량서류는 거부가 정답.",
    "  - (정답·적발) = 수배자/테러 등 특수 적발 분기 대사.",
    "  - 모션/UI 연출 행과 오거부횟수는 제외(대사만).",
]
for i, t in enumerate(lines, start=2):
    info[f"A{i}"] = t
info.column_dimensions["A"].width = 100

wb.save(OUT)
print("saved:", OUT)
print("visitors:", visitor_idx)
