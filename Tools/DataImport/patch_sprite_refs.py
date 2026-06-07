# -*- coding: utf-8 -*-
"""
patch_sprite_refs.py — 「여권 주세요」 사진 참조를 한글 이름(실존 PNG)으로 정규화 (data-tools)

배경: customer.sprite_ref / passport.photo_ref 가 코드값(spr_cNNN / photo_N)이라
      Resources/Characters/<한글이름>.png 와 안 맞아 초상/여권 사진이 로드 실패(플레이스홀더)했다.
      런타임(CustomerView/DocumentCardView)은 Resources.Load("Characters/"+StripExt(spriteRef))
      로 로드하므로 spriteRef 가 손님 한글 이름이면 실제 사진이 뜬다.

이 스크립트가 메인 엑셀(data/여권_정리_updated.xlsx)을 in-place 로 정규화한다(idempotent):
  1) customer.sprite_ref  = 그 손님 name_kr (NAME_OVERRIDE 로 에셋명 보정)
  2) passport.photo_ref   = 그 손님 sprite 이름(정상 손님은 초상과 동일 → 얼굴 대조 일치)
  3) fake_value_pool       = face 디코이(실존 PNG 이름) 행을 표준 셋으로 재생성
                             (사진 결함 MISMATCH_PHOTO 가 본인 아닌 실존 인물로 변조되게)

사진 결함 자체(MISMATCH_PHOTO)는 런타임 변조라 여기서 주입하지 않는다 —
build_days.apply_defect 가 face 풀에서 본인 제외 디코이를 결정론적으로 끼운다(진실만 저장 원칙).

NAME_OVERRIDE: 엑셀 name_kr 과 실존 에셋 파일명이 다른 경우 보정.
  - 35 '자오 레이' → 실존 에셋 '지오 레이'(폴더에 자오 레이.png 없음, 지오 레이.png 존재).

실행: python Tools/DataImport/patch_sprite_refs.py
주의: 엑셀이 Excel 에서 열려 있으면 PermissionError → Excel 을 닫고 재실행.
"""
import os
import sys
import unicodedata

try:
    import openpyxl
except ImportError:
    sys.stderr.write("openpyxl 필요: pip install openpyxl\n")
    sys.exit(1)

_HERE = os.path.dirname(os.path.abspath(__file__))
_REPO = os.path.abspath(os.path.join(_HERE, "..", ".."))
XLSX = os.path.join(_REPO, "data", "여권_정리_updated.xlsx")
CHARDIR = os.path.join(_REPO, "Assets", "Resources", "Characters")

# 엑셀 name_kr → 실존 에셋명 보정(파일명이 다른 손님만).
NAME_OVERRIDE = {"35": "지오 레이"}

# 사진 불일치 디코이 풀(전원 실존 PNG). build_days 가 본인 이름을 제외하고 고른다.
# 미배정 실존 에셋 + 타 한국 여성(사진 결함 손님 10/12/21/30 이 전원 한국 여성).
FACE_DECOYS = [
    ("이하민", "사진 불일치 디코이(미배정 실존 에셋)"),
    ("장지원", "사진 불일치 디코이(미배정 실존 에셋)"),
    ("이지은", "사진 불일치 디코이(타 한국 여성)"),
    ("최서연", "사진 불일치 디코이(타 한국 여성)"),
    ("송하늘", "사진 불일치 디코이(타 한국 여성)"),
    ("배수정", "사진 불일치 디코이(타 한국 여성)"),
]


def asset_names():
    out = set()
    for f in os.listdir(CHARDIR):
        if f.lower().endswith(".png"):
            out.add(unicodedata.normalize("NFC", f[:-4]))
    return out


def col_idx_A(ws, eng_name):
    """LAYOUT_A 시트의 row3(영문 헤더)에서 컬럼 인덱스(1-based) 찾기."""
    for c in range(1, ws.max_column + 1):
        v = ws.cell(3, c).value
        if v and str(v).strip() == eng_name:
            return c
    raise KeyError(eng_name)


def norm_id(v):
    if v is None:
        return None
    return str(int(v)) if isinstance(v, (int, float)) else str(v).strip()


def main():
    if not os.path.exists(XLSX):
        sys.stderr.write("입력 없음: %s\n" % XLSX)
        sys.exit(2)
    pngs = asset_names()
    wb = openpyxl.load_workbook(XLSX)

    # 1) customer.sprite_ref
    cust = wb["customer"]
    c_id = col_idx_A(cust, "customer_id")
    c_nk = col_idx_A(cust, "name_kr")
    c_sr = col_idx_A(cust, "sprite_ref")
    namekr_by_id = {}
    changed_cust = 0
    for r in range(5, cust.max_row + 1):
        cid = norm_id(cust.cell(r, c_id).value)
        if cid is None:
            continue
        nk = cust.cell(r, c_nk).value
        sprite = NAME_OVERRIDE.get(cid, str(nk).strip())
        namekr_by_id[cid] = sprite
        if unicodedata.normalize("NFC", sprite) not in pngs:
            sys.stderr.write("경고: 손님 %s 스프라이트 '%s' 에셋 없음\n" % (cid, sprite))
        if cust.cell(r, c_sr).value != sprite:
            cust.cell(r, c_sr).value = sprite
            changed_cust += 1

    # 2) passport.photo_ref = 그 손님 sprite 이름
    pp = wb["passport"]
    p_cid = col_idx_A(pp, "customer_id")
    p_pr = col_idx_A(pp, "photo_ref")
    changed_pp = 0
    for r in range(5, pp.max_row + 1):
        cid = norm_id(pp.cell(r, p_cid).value)
        if cid is None:
            continue
        if cid not in namekr_by_id:
            sys.stderr.write("경고: 여권 행 %d 손님 %s 매칭 안 됨\n" % (r, cid))
            continue
        nm = namekr_by_id[cid]
        if pp.cell(r, p_pr).value != nm:
            pp.cell(r, p_pr).value = nm
            changed_pp += 1

    # 3) fake_value_pool: face 행 재생성(idempotent)
    for name, _ in FACE_DECOYS:
        if unicodedata.normalize("NFC", name) not in pngs:
            sys.stderr.write("경고: 디코이 '%s' 에셋 없음\n" % name)
    fvp = wb["fake_value_pool"]
    ncol = fvp.max_column
    body = []
    for r in range(2, fvp.max_row + 1):
        vals = [fvp.cell(r, c).value for c in range(1, ncol + 1)]
        if all(v is None for v in vals):
            continue
        field = str(vals[1]).strip() if vals[1] is not None else ""
        if field == "face":
            continue  # 기존 face 행 제거(재생성)
        body.append(vals)
    next_id = 1
    for vals in body:
        vals[0] = next_id
        next_id += 1
    for name, note in FACE_DECOYS:
        row = [next_id, "face", name, note] + [None] * (ncol - 4)
        body.append(row[:ncol])
        next_id += 1
    for r in range(fvp.max_row, 1, -1):
        for c in range(1, ncol + 1):
            fvp.cell(r, c).value = None
    for i, vals in enumerate(body):
        for c, v in enumerate(vals, start=1):
            fvp.cell(2 + i, c).value = v

    wb.save(XLSX)
    sys.stdout.write("OK customer=%d passport=%d face_decoys=%d\n"
                     % (changed_cust, changed_pp, len(FACE_DECOYS)))


if __name__ == "__main__":
    main()
