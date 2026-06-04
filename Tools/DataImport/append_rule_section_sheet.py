# -*- coding: utf-8 -*-
"""
append_rule_section_sheet.py — 메인 DB 엑셀에 rule_section 시트를 편입(신설).
(data-tools 소유. SHARED-CONVENTIONS 관계형 레이아웃 LAYOUT_A 4행 헤더)

목적:
  세분화된 규정집(정적 매뉴얼) UI 전용 데이터 테이블 `rule_section`을 신설한다.
  탭(기본/서류/물품) → 좌측 항목 버튼(category) → 우측 내용(title/body/image)을
  데이터로 표현한다. 기존 rule_book(날짜별 브리핑)·dayN.json rules·inspection_notice 는
  전혀 건드리지 않는다(별개 개념).

형식(기존 시트와 동일한 관계형 4행 헤더 = LAYOUT_A):
  row1 = PK/FK 표기(키 컬럼에만)
  row2 = 자료형(int / varchar(N) / text)
  row3 = 영문 컬럼명(snake_case, 매핑 기준)
  row4 = 한글 라벨
  row5+ = 데이터

스키마(컬럼):
  section_id(int PK), tab(varchar20), tab_order(int), category(varchar30),
  category_order(int), title(varchar60), body(text), image_ref(varchar60),
  content_type(varchar10), unlock_day(int), priority(int),
  related_field(varchar30), note(varchar120)

idempotent:
  - 실행 전 타임스탬프 백업.
  - 기존 시트는 읽지도 쓰지도 않는다.
  - rule_section 시트가 이미 있으면 제거 후 재생성(같은 입력 → 같은 결과).
  - 맨 뒤에 append.

사용:
    python append_rule_section_sheet.py [메인.xlsx]
인자 없으면 기본 경로.

콘솔이 cp949라 한글 stdout이 깨진다 → 결과 확인은 생성된 xlsx를 openpyxl로 Read.
"""
import sys
import os
import shutil
import datetime

try:
    import openpyxl
except ImportError:
    sys.stderr.write("openpyxl 필요: pip install openpyxl\n")
    sys.exit(1)

_REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
DEFAULT_XLSX = os.path.join(_REPO, "data", "여권_정리_updated.xlsx")

SHEET_NAME = "rule_section"

# ── 컬럼 정의(pkfk / type / 영문키 / 한글 라벨) — 단일 소스 ──────────────
COLS = [
    {"key": "section_id",     "type": "int",          "label": "섹션ID",      "pkfk": "PK"},
    {"key": "tab",            "type": "varchar(20)",  "label": "탭",          "pkfk": ""},
    {"key": "tab_order",      "type": "int",          "label": "탭순서",      "pkfk": ""},
    {"key": "category",       "type": "varchar(30)",  "label": "항목",        "pkfk": ""},
    {"key": "category_order", "type": "int",          "label": "항목순서",    "pkfk": ""},
    {"key": "title",          "type": "varchar(60)",  "label": "제목",        "pkfk": ""},
    {"key": "body",           "type": "text",         "label": "본문",        "pkfk": ""},
    {"key": "image_ref",      "type": "varchar(60)",  "label": "이미지키",    "pkfk": ""},
    {"key": "content_type",   "type": "varchar(10)",  "label": "내용유형",    "pkfk": ""},
    {"key": "unlock_day",     "type": "int",          "label": "노출일자",    "pkfk": ""},
    {"key": "priority",       "type": "int",          "label": "우선순위",    "pkfk": ""},
    {"key": "related_field",  "type": "varchar(30)",  "label": "연계필드",    "pkfk": ""},
    {"key": "note",           "type": "varchar(120)", "label": "비고",        "pkfk": ""},
]

# ── 데이터 6행 (확정 내용) ─────────────────────────────────────────────
#  탭 묶음: 기본=기본규정·특별지시 / 서류=신분·유효기간·국적 / 물품=금지물품
#  tab_order: 기본1 서류2 물품3. category_order: 각 탭 내 1,2,...
ROWS = [
    {
        "section_id": 1, "tab": "기본", "tab_order": 1,
        "category": "기본 규정", "category_order": 1,
        "title": "기본 입국 규정",
        "body": "모든 입국자는 유효한 여권을 제시해야 한다. 미제출·서류 결함 시 거부.",
        "image_ref": "rb_passport", "content_type": "both",
        "unlock_day": 1, "priority": 0, "related_field": "여권", "note": "",
    },
    {
        "section_id": 2, "tab": "기본", "tab_order": 1,
        "category": "특별 지시", "category_order": 2,
        "title": "특별 지시(당일 브리핑)",
        "body": "뉴스·브리핑으로 내려오는 당일 지시는 기존 규정보다 우선 적용한다(예: 특정 국적 제한, 검역 강화). 그날의 지시는 규정집이 아니라 브리핑을 확인할 것.",
        "image_ref": "", "content_type": "text",
        "unlock_day": 1, "priority": 10, "related_field": "",
        "note": "당일 브리핑(dayN rules) 우선",
    },
    {
        "section_id": 3, "tab": "서류", "tab_order": 2,
        "category": "신분 확인", "category_order": 1,
        "title": "신분(본인) 확인",
        "body": "여권의 사진·성별·생년월일이 본인과 일치해야 한다. 사진과 외모가 다르면 지문 인식으로 확인.",
        "image_ref": "rb_identity", "content_type": "both",
        "unlock_day": 1, "priority": 0, "related_field": "사진", "note": "",
    },
    {
        "section_id": 4, "tab": "서류", "tab_order": 2,
        "category": "여권 유효기간", "category_order": 2,
        "title": "여권 유효기간",
        "body": "입국일 기준 만료일이 지난 여권은 입국 불가. 발급일이 미래인 여권은 위조 의심.",
        "image_ref": "", "content_type": "text",
        "unlock_day": 4, "priority": 0, "related_field": "유효기간", "note": "",
    },
    {
        "section_id": 5, "tab": "서류", "tab_order": 2,
        "category": "국적별 조건", "category_order": 3,
        "title": "국적별 입국 조건",
        "body": "외국 국적자는 유효한 비자 필수(3일차~). 여권 발급국·국적·비자가 일치해야 한다. 여권번호는 발급국 형식과 일치. 장기체류자는 취업/합격 증명 필요.",
        "image_ref": "rb_visa", "content_type": "both",
        "unlock_day": 3, "priority": 0, "related_field": "국적", "note": "",
    },
    {
        "section_id": 6, "tab": "물품", "tab_order": 3,
        "category": "금지 물품", "category_order": 1,
        "title": "반입 금지 물품",
        "body": "마약·밀수품·금괴 등 반입 불가. X-ray 정밀 검사로 적발. 금지 물품 수령 시 공범으로 간주(즉시 처벌).",
        "image_ref": "rb_contraband", "content_type": "both",
        "unlock_day": 11, "priority": 0, "related_field": "물품", "note": "",
    },
]


def write_sheet(ws, cols, rows):
    """관계형 4행 헤더 + 데이터. row1 pkfk / row2 type / row3 eng / row4 kr / row5+ data."""
    for ci, col in enumerate(cols, start=1):
        pkfk = (col.get("pkfk") or "").strip()
        ws.cell(row=1, column=ci, value=pkfk if pkfk else None)
        ws.cell(row=2, column=ci, value=col["type"])
        ws.cell(row=3, column=ci, value=col["key"])
        ws.cell(row=4, column=ci, value=col["label"])
    keys = [c["key"] for c in cols]
    for ri, row in enumerate(rows, start=5):
        for ci, k in enumerate(keys, start=1):
            v = row.get(k)
            ws.cell(row=ri, column=ci, value=v if v not in ("", None) else None)


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    if not os.path.exists(xlsx):
        sys.stderr.write("메인 엑셀 없음: %s\n" % xlsx)
        sys.exit(2)

    # ── 백업(타임스탬프) ──────────────────────────────────────
    ts = datetime.datetime.now().strftime("%y%m%d_%H%M%S")
    base, ext = os.path.splitext(xlsx)
    backup = "%s.backup_%s%s" % (base, ts, ext)
    shutil.copy2(xlsx, backup)

    wb = openpyxl.load_workbook(xlsx)  # 식 보존
    existing = list(wb.sheetnames)

    # ── idempotent: 이미 있으면 제거 후 재생성 ────────────────
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]

    ws = wb.create_sheet(title=SHEET_NAME)
    write_sheet(ws, COLS, ROWS)
    wb.save(xlsx)

    # ── 요약(UTF-8 파일) ─────────────────────────────────────
    summ = []
    summ.append("backup: %s" % backup)
    summ.append("main xlsx: %s" % xlsx)
    summ.append("기존 시트 수: %d" % len(existing))
    summ.append("최종 시트 수: %d" % len(wb.sheetnames))
    summ.append("최종 시트 순서: %s" % ", ".join(wb.sheetnames))
    summ.append("추가됨: %s rows=%d" % (SHEET_NAME, len(ROWS)))
    summ_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_append_rule_section_summary.txt")
    with open(summ_path, "w", encoding="utf-8") as f:
        f.write("\n".join(summ))

    sys.stdout.write("OK sheets=%d backup=%s\n" % (len(wb.sheetnames), os.path.basename(backup)))


if __name__ == "__main__":
    main()
