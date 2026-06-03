# -*- coding: utf-8 -*-
"""dialogue_case 스키마 개편:
 - character_type 컬럼 재추가 (customer 시트에서 채움)
 - game_result -> verdict + result 2컬럼 분리
 - document_type: violation_field='없음' 행은 '-'
 - case 8 violation_field '비자' -> '만료일' 수정
새 컬럼 순서:
 dialogue_case_id, schedule_id, customer_id, character_type, case_type,
 document_type, violation_field, verdict, result, reject_count
"""
import shutil
from openpyxl import load_workbook

p = 'data/여권_정리_updated.xlsx'
bak = 'data/여권_정리_updated.backup_before_verdict_split.xlsx'
shutil.copyfile(p, bak)
print('백업:', bak)

wb = load_workbook(p)

# customer_id -> character_type
cu = wb['customer']
uk = [c.value for c in cu[3]]
u_id = uk.index('customer_id'); u_ct = uk.index('character_type')
cust_type = {}
for r in range(5, cu.max_row + 1):
    cid = cu.cell(row=r, column=u_id + 1).value
    if cid is not None:
        cust_type[cid] = cu.cell(row=r, column=u_ct + 1).value

dc = wb['dialogue_case']
k = [c.value for c in dc[3]]
idx = {name: k.index(name) for name in k}

GR_MAP = {
    '정상 승인': ('입국 허가', '정답'),
    '정상 거절': ('입국 거부', '정답'),
    '잘못 허가': ('입국 허가', '오답'),
    '잘못 거절': ('입국 거부', '오답'),
    '-': ('-', '-'),
}

# 기존 데이터 읽기
data = []
for r in range(5, dc.max_row + 1):
    row = [dc.cell(row=r, column=c + 1).value for c in range(len(k))]
    if all(v in (None, '') for v in row):
        continue
    rec = dict(zip(k, row))
    data.append(rec)

new_rows = []
for rec in data:
    cid = rec['dialogue_case_id']
    vf = rec['violation_field']
    # case 8 오류 수정
    if cid == 8 and vf == '비자':
        vf = '만료일'
    # document_type: 위반 없으면 '-'
    doc = rec['document_type']
    doc = '-' if vf == '없음' else doc
    # verdict/result 분리
    gr = rec['game_result']
    verdict, result = GR_MAP.get(gr, ('-', '-'))
    new_rows.append([
        rec['dialogue_case_id'], rec['schedule_id'], rec['customer_id'],
        cust_type.get(rec['customer_id'], ''), rec['case_type'],
        doc, vf, verdict, result, rec['reject_count'],
    ])

# 새 헤더(4줄)
H1 = ['PK', 'FK', 'FK', '', '', '', '', '', '', '']
H2 = ['int', 'int', 'int', 'varchar(30)', 'varchar(30)', 'varchar(20)', 'varchar(30)', 'varchar(20)', 'varchar(10)', 'int']
H3 = ['dialogue_case_id', 'schedule_id', 'customer_id', 'character_type', 'case_type', 'document_type', 'violation_field', 'verdict', 'result', 'reject_count']
H4 = ['대화 흐름 ID', '일정 ID', '고객 ID', '캐릭터 유형', '상황 유형', '문서 종류', '위반 항목', '판정', '정오', '거절 횟수']

# 시트 재생성(같은 위치)
pos = wb.sheetnames.index('dialogue_case')
del wb['dialogue_case']
ws = wb.create_sheet('dialogue_case', pos)
for row in (H1, H2, H3, H4):
    ws.append(row)
for row in new_rows:
    ws.append(row)

wb.save(p)
print('저장 완료. 데이터 행수:', len(new_rows))

# 검증
wb2 = load_workbook(p, read_only=True)
d2 = wb2['dialogue_case']
print('새 헤더:', [c.value for c in next(d2.iter_rows(min_row=3, max_row=3))])
import collections
k2 = [c.value for c in next(d2.iter_rows(min_row=3, max_row=3))]
vi = k2.index('verdict'); ri = k2.index('result'); di = k2.index('document_type'); ci = k2.index('character_type')
vc = collections.Counter(); rc = collections.Counter(); dcd = collections.Counter()
for r in d2.iter_rows(min_row=5, values_only=True):
    if all(c in (None, '') for c in r):
        continue
    vc[r[vi]] += 1; rc[r[ri]] += 1; dcd[r[di]] += 1
print('verdict 분포:', dict(vc))
print('result 분포:', dict(rc))
print('document_type 분포:', dict(dcd))
print('샘플 3행:')
for r in list(d2.iter_rows(min_row=5, max_row=7, values_only=True)):
    print('   ', r)
