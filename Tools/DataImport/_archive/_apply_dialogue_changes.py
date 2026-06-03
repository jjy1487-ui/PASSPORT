# -*- coding: utf-8 -*-
import shutil
from openpyxl import load_workbook

p = 'data/여권_정리_updated.xlsx'
bak = 'data/여권_정리_updated.backup_before_speaker_id.xlsx'
shutil.copyfile(p, bak)
print('백업:', bak)

wb = load_workbook(p)
dc = wb['dialogue_case']
dl = wb['dialogue_line']

# ── 맵 만들기 ──────────────────────────────────────────
ck = [c.value for c in dc[3]]
dc_cid = ck.index('customer_id')
dc_id = ck.index('dialogue_case_id')
case2cust = {}
for r in range(5, dc.max_row + 1):
    cid = dc.cell(row=r, column=dc_id + 1).value
    if cid is None:
        continue
    case2cust[cid] = dc.cell(row=r, column=dc_cid + 1).value

cu = wb['customer']
uk = [c.value for c in cu[3]]
u_id = uk.index('customer_id'); u_nm = uk.index('name_kr')
cust2name = {}
for r in range(5, cu.max_row + 1):
    cid = cu.cell(row=r, column=u_id + 1).value
    if cid is None:
        continue
    cust2name[cid] = cu.cell(row=r, column=u_nm + 1).value

# ── 1) dialogue_line 에 speaker_id 컬럼 삽입 (speaker=4열 뒤 → 5열) ──
dlk = [c.value for c in dl[3]]
sp_col = dlk.index('speaker') + 1          # 4
new_col = sp_col + 1                        # 5
dl.insert_cols(new_col)
dl.cell(row=1, column=new_col, value='FK')
dl.cell(row=2, column=new_col, value='int')
dl.cell(row=3, column=new_col, value='speaker_id')
dl.cell(row=4, column=new_col, value='화자 캐릭터 ID')

# ── 2) 기존 데이터: 캐릭터→이름, speaker_id 채우기 ──
dcid_col = dlk.index('dialogue_case_id') + 1
changed = 0
for r in range(5, dl.max_row + 1):
    sp = dl.cell(row=r, column=sp_col).value
    if sp is None:
        continue
    caseid = dl.cell(row=r, column=dcid_col).value
    cust = case2cust.get(caseid)
    if sp == '캐릭터':
        nm = cust2name.get(cust)
        if nm:
            dl.cell(row=r, column=sp_col, value=nm)
        dl.cell(row=r, column=new_col, value=cust)
        changed += 1
    elif sp == '심사관':
        dl.cell(row=r, column=new_col, value=None)  # 심사관은 ID 없음
print('기존 캐릭터→이름 변경:', changed, '줄')

# ── 3) Part B: 특수 캐릭터 다단계 항의 추가 ──
maxcase = max(case2cust.keys())
# 현재 max line_id
li_col = dlk.index('line_id') + 1
maxline = 0
for r in range(5, dl.max_row + 1):
    v = dl.cell(row=r, column=li_col).value
    if isinstance(v, int):
        maxline = max(maxline, v)

# (customer_id, schedule_id, name, [ (rc, [(is_inspector, text), ...]) ... ])
specials = [
 (15, 30, '한지원', [
   (1, [(True, '죄송하지만 입국이 어렵습니다.'),
        (False, '네? 저 한지원인데요. 뭔가 착오가 있는 거 아니에요?')]),
   (2, [(True, '다시 확인했지만 입국이 어렵습니다.'),
        (False, '이거 SNS에 올라가면 곤란해지실 텐데요. 다시 봐주세요.')]),
   (3, [(True, '죄송합니다만 지나갈 수 없습니다.'),
        (False, '지금 저 막은 거 후회하실 거예요. 매니저한테 전화할게요!'),
        (True, '…다시 확인하니 서류에 문제가 없으십니다. 통과하셔도 됩니다.'),
        (False, '진작 좀 알아보지 그랬어요. 흥.')]),
 ]),
 (16, 54, '오현석', [
   (1, [(True, '죄송하지만 입국이 어렵습니다.'),
        (False, '허허, 무엇이 문제인지 다시 한번 살펴보시게.')]),
   (2, [(True, '다시 확인했지만 입국이 어렵습니다.'),
        (False, '급히 보면 보이던 것도 안 보이는 법. 차분히 보시게나.')]),
   (3, [(True, '죄송합니다만 지나갈 수 없습니다.'),
        (False, '사람은 누구나 실수하지. 허나 같은 실수를 거듭해선 안 되는 법일세.'),
        (True, '…다시 보니 이상이 없으십니다. 통과하셔도 됩니다.'),
        (False, '괜찮네. 다음부턴 마음을 가라앉히고 보시게.')]),
 ]),
 (17, 60, '윤정호', [
   (1, [(True, '죄송하지만 입국이 어렵습니다.'),
        (False, '이보시오, 내가 누군지 알고 이러시오? 다시 확인해 보시오.')]),
   (2, [(True, '다시 확인했지만 입국이 어렵습니다.'),
        (False, '허, 이 사람 보게. 윗선에 한마디면 그 자리 위태로울 텐데?')]),
   (3, [(True, '죄송합니다만 지나갈 수 없습니다.'),
        (False, '내 이 일을 그냥 넘기지 않겠소! 당장 책임자 부르시오!'),
        (True, '…다시 확인하니 문제가 없으십니다. 통과하셔도 됩니다.'),
        (False, '진작 그럴 것이지. 앞으로 똑바로 일하시오.')]),
 ]),
]

ncase = maxcase
nline = maxline
added_cases = 0; added_lines = 0
for cust_id, sched_id, name, stages in specials:
    for rc, lines in stages:
        ncase += 1
        # dialogue_case 행: [case_id, schedule_id, customer_id, case_type, document_type, violation_field, game_result, reject_count]
        dc.append([ncase, sched_id, cust_id, '일반 심사', '여권', '없음', '잘못 거절', rc])
        added_cases += 1
        for order, (is_insp, text) in enumerate(lines, 1):
            nline += 1
            spk = '심사관' if is_insp else name
            sid = None if is_insp else cust_id
            # dialogue_line 행: [line_id, dialogue_case_id, line_order, speaker, speaker_id, text_kr]
            dl.append([nline, ncase, order, spk, sid, text])
            added_lines += 1
print('Part B 추가 → case:', added_cases, ', line:', added_lines)

wb.save(p)
print('저장 완료')

# ── 검증 ──
wb2 = load_workbook(p, read_only=True)
dl2 = wb2['dialogue_line']
print('dialogue_line 헤더:', [c.value for c in next(dl2.iter_rows(min_row=3, max_row=3))])
print('샘플(데이터 2행):')
for r in list(dl2.iter_rows(min_row=5, max_row=6, values_only=True)):
    print('   ', r)
print('한지원 줄 확인:')
for r in dl2.iter_rows(min_row=5, values_only=True):
    if r[3] == '한지원':
        print('   ', r)
