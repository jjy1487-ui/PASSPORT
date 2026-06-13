# -*- coding: utf-8 -*-
"""손글 원본(방문객_스크립트_일자별_생성_손글기반_260611.xlsx)을 파싱해
시나리오_스크립트.xlsx 의 '대사_스크립트' 시트를 재작성한다.
- 확률 캐릭터(박철수 등)는 손글에 [서류 정상]/[서류 불량] 블록이 둘 다 있으므로 양쪽 모두 표시.
- 모션/UI/판정/검수 행 제외, 대사(인삿말·검문관·방문객반응)만.
- day5: 한지원(2번째)↔장우진(5번째) 순서 스왑 적용(사용자 요청).
파일이 Excel에 열려 있어도 xlwings로 그 인스턴스에 붙어 갱신.
"""
import os, re, io
import openpyxl, xlwings as xw

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HAND = r"C:\Users\chris\Downloads\방문객_스크립트_일자별_생성_손글기반_260611.xlsx"
OUTNAME = "시나리오_스크립트.xlsx"
OUTPATH = os.path.join(ROOT, "data", OUTNAME)

DIALOGUE_STEPS = {"인삿말", "검문관", "방문객반응"}

def clean(h):
    if h is None: return None
    s = str(h)
    m = re.search(r"「(.*?)」", s, re.S)
    if m: return m.group(1).strip()
    return s.split("※")[0].strip().strip('"').strip()

def parse_header(text):
    # "[1번째 방문객] 김민준  |  일반  |  서류: 정상  (더미·클론←X)"
    m = re.search(r"\[(\d+)\s*번째\s*방문객\]\s*([^|]+?)\s*\|\s*([^|]+?)\s*\|\s*서류:\s*([^|]+)", text)
    if not m: return None
    order = int(m.group(1)); name = m.group(2).strip(); typ = m.group(3).strip()
    docstate = m.group(4).strip()
    docstate = re.sub(r"\(더미.*?\)", "", docstate).strip()
    return order, name, typ, docstate

wb = openpyxl.load_workbook(HAND, data_only=True)
blocks = []  # dict: day, order, name, typ, docstate, intro[], approve[], reject[]
for sheet in [s for s in wb.sheetnames if s.startswith("Day")]:
    day = int(re.search(r"(\d+)", sheet).group(1))
    ws = wb[sheet]
    cur = None
    for r in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c) for c in r]
        joined = " ".join(cells)
        if "번째 방문객]" in joined:
            hdr = next((c for c in cells if "번째 방문객]" in c), joined)
            p = parse_header(hdr)
            if p:
                cur = {"day": day, "order": p[0], "name": p[1], "typ": p[2],
                       "docstate": p[3], "intro": [], "approve": [], "reject": []}
                blocks.append(cur)
            continue
        if cur is None: continue
        step = cells[1].strip() if len(cells) > 1 else ""
        if step not in DIALOGUE_STEPS: continue
        result = cells[6].strip() if len(cells) > 6 else ""
        name = cells[3].strip() if len(cells) > 3 else ""
        line = clean(cells[7] if len(cells) > 7 else "")
        if not line: continue
        speaker = "심사관" if step == "검문관" else (name or cur["name"])
        entry = (speaker, line)
        if step == "인삿말" or result == "공통":
            cur["intro"].append(entry)
        elif result == "허가":
            cur["approve"].append(entry)
        elif result == "거부":
            cur["reject"].append(entry)

# day5 swap: 한지원 <-> 장우진 순서
for b in blocks:
    if b["day"] == 5:
        if b["name"] == "한지원": b["order"] = 5
        elif b["name"] == "장우진": b["order"] = 2
# stable sort by (day, order); 같은 order(정상/불량)는 등장 순서 유지
blocks.sort(key=lambda b: (b["day"], b["order"]))

# build rows
header = ["일차","순서","방문객","유형","서류상태","분기","화자","대사"]
rows = [header]
name_rows = []
for b in blocks:
    normal = ("정상" in b["docstate"]) and ("불량" not in b["docstate"])
    approve_label = "입국 허가 (정답)" if normal else "입국 허가 (오판)"
    reject_label = "입국 거부 (오판)" if normal else "입국 거부 (정답)"
    sections = [("인삿말", b["intro"]), (approve_label, b["approve"]), (reject_label, b["reject"])]
    first = True
    for label, lines in sections:
        for i, (spk, txt) in enumerate(lines):
            head = (label == "인삿말" and i == 0)
            rows.append([b["day"] if head else "", b["order"] if head else "",
                         b["name"] if head else "", b["typ"] if head else "",
                         b["docstate"] if head else "", label if i == 0 else "", spk, txt])
            if first:
                name_rows.append(len(rows)); first = False
    rows.append([""]*8)

# write via xlwings (attach if open)
bk = None
for app in xw.apps:
    for b in app.books:
        try:
            if os.path.basename(b.fullname).lower() == OUTNAME.lower(): bk = b; break
        except Exception: pass
    if bk: break
opened = False
if bk is None:
    app = xw.App(visible=False); app.display_alerts = False; bk = app.books.open(OUTPATH); opened = True
ws = bk.sheets["대사_스크립트"]
ws.clear()
ws.range((1,1)).value = rows
hdr = ws.range((1,1),(1,8)); hdr.color = (47,85,151); hdr.api.Font.Bold = True; hdr.api.Font.Color = 0xFFFFFF
for col, wdt in zip("ABCDEFGH", [6,6,14,16,28,18,10,70]):
    ws.range(f"{col}1").column_width = wdt
ws.range((2,8),(len(rows),8)).api.WrapText = True
for r in name_rows: ws.range((r,3)).api.Font.Bold = True
# branch colors
GREEN = 31 + 122*256 + 31*65536; RED = 176 + 36*256 + 24*65536
vals = ws.range((2,6),(len(rows),6)).value
if not isinstance(vals, list): vals = [vals]
for i, v in enumerate(vals):
    if not v: continue
    rr = 2 + i
    if "허가" in str(v): ws.range((rr,6)).api.Font.Color = GREEN; ws.range((rr,6)).api.Font.Bold = True
    elif "거부" in str(v): ws.range((rr,6)).api.Font.Color = RED; ws.range((rr,6)).api.Font.Bold = True
bk.save()
if opened: bk.app.quit()

# verification dump
dual = {}
for b in blocks:
    dual.setdefault((b["day"], b["name"]), []).append(b["docstate"])
multi = {k: v for k, v in dual.items() if len(v) > 1}
print(f"blocks={len(blocks)} visitors_distinct={len(dual)} dual_variant_chars={len(multi)} rows={len(rows)}")
print("dual-variant examples:", list(multi.items())[:6])
