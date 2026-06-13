# -*- coding: utf-8 -*-
"""대사_스크립트(마스터)를 일차별 시트(Day1..Day14)로 분리한 파일 생성."""
import openpyxl, re, sys
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
sys.stdout.reconfigure(encoding='utf-8')
SRC='data/시나리오_스크립트.xlsx'; OUT='data/대사_스크립트.xlsx'
DAYS={
 1:("첫 번째 도장","신입 관리직 배치 · 기본 심사 튜토리얼"),
 2:("여기가 BTS의 나라입니까?","외국인 관광객 증가"),
 3:("강남은 어디에 있나요?","성형 수술 고객 입국 · 지문·비자"),
 4:("잠깐만요, 다시 보여주세요","입국 심사 불만 증가 · 방문목적 확인"),
 5:("검역 강화","방역 체계 가동 · PCR 검사서"),
 6:("긴장 속의 일상","방역 체제 유지"),
 7:("다음 사람, 다음 사람","피로 누적"),
 8:("새로운 시작","유학생 입국 · 입학 합격서"),
 9:("꿈을 찾아서","취업 목적 방문객 · 재직 증명서"),
 10:("평범한 하루","무난한 하루"),
 11:("그 사람을 왜 통과시켰죠?","범죄 조직 유입 · X-RAY"),
 12:("서울 테러 속보","보안 경계 격상 · 범죄자 등장"),
 13:("정말 괜찮은 사람일까?","의심의 일상화"),
 14:("최종 심사","마지막 근무 · 단일 엔딩"),
}
COLS=["일차","순서","방문객","유형","서류상태","분기","화자","대사"]
# ── 마스터 읽기 → 일차별 행 묶기 ──
wb=openpyxl.load_workbook(SRC,data_only=True)
ws=wb['대사_스크립트']
rows=[list(r) for r in ws.iter_rows(values_only=True)]
def s(v): return '' if v is None else str(v)
buckets={d:[] for d in DAYS}; cur=None
for r in rows[1:]:
    a=s(r[0])
    m=re.match(r'━+\s*(\d+)일차',a)
    if m: cur=int(m.group(1)); continue   # 배너 → 일차 전환, 배너는 버림
    if all(s(x).strip()=='' for x in r[:8]):  # 빈 줄
        if cur: buckets[cur].append(['']*8); continue
    if cur is None:
        d2=None
        if a.strip().replace('.0','').isdigit(): d2=int(float(a))
        if d2 in DAYS: cur=d2
    if cur: buckets[cur].append([s(x) for x in r[:8]])

# ── 출력 워크북: Day1..Day14 ──
out=openpyxl.Workbook(); out.remove(out.active)
HBANNER=PatternFill('solid',fgColor='1F4E97'); HCOL=PatternFill('solid',fgColor='2F5597')
WHITE=Font(bold=True,color='FFFFFF'); WHITE12=Font(bold=True,color='FFFFFF',size=12)
GREEN=Font(bold=True,color='1F7A1F'); RED=Font(bold=True,color='B02418'); PURPLE=Font(bold=True,color='7030A0')
NFILL=PatternFill('solid',fgColor='E2EFDA'); DFILL=PatternFill('solid',fgColor='FCE4D6')
WRAP=Alignment(wrap_text=True,vertical='top'); TOP=Alignment(vertical='top')
thin=Side(style='thin',color='D9D9D9'); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
CROSS={'지문 검사','지문 대조','서류 대조','음성 대조','경보 대조','X-ray 검사'}
W=[5,5,13,16,26,16,10,60]
for d in DAYS:
    title,theme=DAYS[d]
    sh=out.create_sheet(f'Day{d}')
    # row1 배너
    sh.append([f'{d}일차 · {title}  —  {theme}']+['']*7)
    for c in range(1,9): sh.cell(1,c).fill=HBANNER
    sh.cell(1,1).font=WHITE12
    # row2 컬럼헤더
    sh.append(COLS)
    for c in range(1,9):
        cc=sh.cell(2,c); cc.fill=HCOL; cc.font=WHITE; cc.alignment=Alignment(horizontal='center',vertical='center'); cc.border=BORDER
    # 본문
    for rv in buckets[d]:
        sh.append(rv)
        rr=sh.max_row
        br=rv[5]; st=rv[4]
        for c in range(1,9):
            cc=sh.cell(rr,c); cc.alignment=(WRAP if c==8 else TOP); cc.border=BORDER
        if br:
            if br.startswith('입국 허가'): sh.cell(rr,6).font=GREEN
            elif br.startswith('입국 거부'): sh.cell(rr,6).font=RED
            elif br in CROSS: sh.cell(rr,6).font=PURPLE
        if rv[2]: sh.cell(rr,3).font=Font(bold=True)   # 방문객명 굵게
        if st:
            sh.cell(rr,5).fill = NFILL if st.startswith('정상') else (DFILL if st.startswith('불량') else PatternFill())
    for i,w in enumerate(W,1): sh.column_dimensions[chr(64+i)].width=w
    sh.freeze_panes='A3'
out.save(OUT)
print('생성:',OUT)
print('시트:',out.sheetnames)
for d in DAYS: print(f'  Day{d}: {len(buckets[d])}행')
