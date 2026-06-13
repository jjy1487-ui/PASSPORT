# -*- coding: utf-8 -*-
"""엑셀 서류상태 칸을 구체 결함 라벨로 갱신 (마스터 + 날짜별 파일)."""
import json,glob,re,os,sys
sys.stdout.reconfigure(encoding='utf-8')

def pidmap(docs):
    pp=next((x for x in (docs or []) if x.get('documentType')=='여권'),None)
    if not pp: return {}
    g=lambda k: next((f['value'] for f in pp.get('fields',[]) if f['key']==k),None)
    return {'name':g('name'),'birth':g('birth_date'),'nat':g('nationality')}
FMAP={('여권','만료일'):'여권 만료일 경과',('여권','생년월일'):'여권 생년월일 불일치',('여권','이름'):'여권 이름 불일치',
 ('여권','사진'):'여권 사진 불일치',('여권','국적'):'여권 국적 불일치',('여권','성별'):'여권 성별 불일치',
 ('PCR검사서','검사 결과'):'PCR 검사결과 양성',('PCR검사서','검사 기관'):'PCR 검사기관 불일치',('PCR검사서','검사일'):'PCR 검사일 만료',
 ('취업증빙','고용 회사'):'취업증빙 고용회사 불일치',('취업증빙','입사일'):'취업증빙 입사일 불일치'}
XMAP={'밀수품':'X-ray 밀수품 적발','마약':'X-ray 마약 적발','폭발물 부품':'X-ray 폭발물 적발'}
def fp_field(var):
    P=pidmap(var.get('documents')); rec=(var.get('fingerprint') or {}).get('record') or {}
    if rec.get('mode')=='수배자' or (rec.get('criminalRecord','') not in ('','없음')): return '지문 수배자 일치'
    if (rec.get('dbName') or '').upper()!=(P.get('name') or '').upper(): return '지문 이름 불일치'
    if rec.get('dbBirth')!=P.get('birth'): return '지문 생년월일 불일치'
    if (rec.get('dbNationality') or '') not in (P.get('nat') or ''): return '지문 국적 불일치'
    return '지문 신원 불일치'
def label(var):
    parts=[]
    for doc in (var.get('documents') or []):
        vf=doc.get('violationField','')
        if doc.get('variant')=='비정상' and vf not in ('','없음'):
            dt=doc['documentType']
            if dt=='여권' and vf=='여권번호': continue
            parts.append(f'비자 {vf} 불일치' if dt=='비자' else FMAP.get((dt,vf),f'{dt} {vf}'))
    xr=var.get('xray')
    if xr and (xr.get('result')=='불일치' or xr.get('detail')): parts.append(XMAP.get(xr.get('detail',''),'X-ray 적발'))
    fp=var.get('fingerprint')
    if fp and fp.get('result')=='불일치': parts.append(fp_field(var))
    if not parts and var.get('defectVariant'): parts.append(var['defectVariant'])
    return ' + '.join(parts) if parts else '결함'

# (day,name) -> '불량(라벨)'
LAB={}
for p in sorted(glob.glob('Assets/Resources/GameData/day*.json'),key=lambda x:int(re.search(r'day(\d+)',x).group(1))):
    if 'backup' in p: continue
    day=int(re.search(r'day(\d+)',p).group(1))
    for c in json.load(open(p,encoding='utf-8')).get('customers',[]):
        for var in [c,c.get('altVariant') or {}]:
            if var and '거절' in (var.get('correctResult') or ''):
                LAB[(day,c['nameKr'])]=f'불량({label(var)})'; break

import xlwings as xw
def is_open(path):
    for app in xw.apps:
        for b in app.books:
            try:
                if os.path.basename(b.fullname).lower()==os.path.basename(path).lower(): return app,b
            except: pass
    return None

def apply_openpyxl(path):
    import openpyxl
    wb=openpyxl.load_workbook(path); n=0
    for sh in wb.sheetnames:
        ws=wb[sh]
        sheetday=None
        m=re.match(r'Day(\d+)',sh)
        if m: sheetday=int(m.group(1))
        for row in ws.iter_rows():
            nm=row[2].value; st=row[4].value
            if nm and st and str(st).strip().startswith('불량'):
                day=sheetday
                if day is None and row[0].value and str(row[0].value).replace('.0','').isdigit(): day=int(float(row[0].value))
                key=(day,str(nm).strip())
                if key in LAB: row[4].value=LAB[key]; n+=1
    wb.save(path); return n

def apply_xlwings(b,sheetnames):
    n=0
    for sh in sheetnames:
        ws=b.sheets[sh]; vals=ws.used_range.value
        sheetday=None
        m=re.match(r'Day(\d+)',sh)
        if m: sheetday=int(m.group(1))
        for i,r in enumerate(vals):
            nm=r[2] if len(r)>2 else None; st=r[4] if len(r)>4 else None
            if nm and st and str(st).strip().startswith('불량'):
                day=sheetday
                a=r[0]
                if day is None and isinstance(a,(int,float)): day=int(a)
                key=(day,str(nm).strip())
                if key in LAB: ws.range((i+1,5)).value=LAB[key]; n+=1
    b.save(); return n

for path,sheets in [('data/대사_스크립트.xlsx',[f'Day{d}' for d in range(1,15)]),
                    ('data/시나리오_스크립트.xlsx',['대사_스크립트'])]:
    full=os.path.abspath(path)
    op=is_open(full)
    if op:
        n=apply_xlwings(op[1], sheets if 'Day1' in [s.name for s in op[1].sheets] else ['대사_스크립트'])
        print(f'{path} (열림/xlwings): {n}칸 갱신')
    else:
        n=apply_openpyxl(full)
        print(f'{path} (닫힘/openpyxl): {n}칸 갱신')
